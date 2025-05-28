from flask import Flask, render_template, request, send_from_directory, jsonify, send_file
from werkzeug.utils import secure_filename
import os
import shutil
from fuzzywuzzy import fuzz
from PyPDF2 import PdfReader
from docx import Document
import openpyxl
import xlrd
import dotenv
import logging
import io
import matplotlib.pyplot as plt
 
dotenv.load_dotenv()
 
app = Flask(__name__)
 
UPLOAD_DIR = "uploaded"
os.makedirs(UPLOAD_DIR, exist_ok=True)
 
app.config['UPLOAD_FOLDER'] = UPLOAD_DIR
 
@app.route("/")
def index():
    return render_template("index.html")
 
@app.route("/dashboard")
def dashboard():
    return render_template("dashboard.html")
 
@app.route("/dashboard/chart.png")
def dashboard_chart():
    plt.figure(figsize=(6,4))
    x = [1, 2, 3, 4, 5]
    y = [10, 20, 15, 30, 25]
    plt.plot(x, y, marker='o')
    plt.title("Sample Chart")
    plt.xlabel("X Axis")
    plt.ylabel("Y Axis")
    plt.tight_layout()
 
    img = io.BytesIO()
    plt.savefig(img, format='png')
    plt.close()
    img.seek(0)
    return send_file(img, mimetype='image/png')
 
@app.route("/chat")
def chat():
    return render_template("chat.html")
@app.route("/overview")
def overview():
    return render_template("overview.html")
 
 
 
@app.route("/upload", methods=["GET", "POST"])
def upload_files():
    if request.method == "POST":
        files = request.files.getlist("files")
        for file in files:
            if file:
                relative_path = file.filename.replace("\\", "/")
                secure_path = os.path.normpath(relative_path)
                dest_path = os.path.join(UPLOAD_DIR, secure_path)
                os.makedirs(os.path.dirname(dest_path), exist_ok=True)
                file.save(dest_path)
        return jsonify({"status": "success", "message": "Files uploaded successfully."})
    return render_template("upload.html")
 
 
@app.route("/reports")
def show_reports():
    all_files = []
    for root, dirs, files in os.walk(UPLOAD_DIR):
        for name in files:
            rel_dir = os.path.relpath(root, UPLOAD_DIR)
            rel_file = os.path.join(rel_dir, name) if rel_dir != '.' else name
            all_files.append(rel_file.replace("\\", "/"))
    return render_template("reports.html", files=all_files)
 
@app.route("/download/<path:filename>")
def download_file(filename):
    full_path = os.path.join(UPLOAD_DIR, filename)
    if os.path.exists(full_path):
        return send_from_directory(UPLOAD_DIR, filename, as_attachment=True)
    return jsonify({"message": "File not found"}), 404
 
def read_txt(file_path):
    try:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()
    except Exception as e:
        return f"Error reading TXT: {str(e)}"
 
def read_pdf(file_path):
    try:
        reader = PdfReader(file_path)
        return "\n".join(page.extract_text() or "" for page in reader.pages)
    except Exception as e:
        return f"Error reading PDF: {str(e)}"
 
def read_docx(file_path):
    try:
        doc = Document(file_path)
        return "\n".join(para.text for para in doc.paragraphs)
    except Exception as e:
        return f"Error reading DOCX: {str(e)}"
 
def read_xlsx(file_path):
    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        all_text = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = [str(cell) for cell in row if cell is not None]
                if row_text:
                    all_text.append(" ".join(row_text))
        return "\n".join(all_text)
    except Exception as e:
        return f"Error reading XLSX: {str(e)}"
 
def read_xls(file_path):
    try:
        wb = xlrd.open_workbook(file_path)
        all_text = []
        for sheet in wb.sheets():
            for row_idx in range(sheet.nrows):
                row = sheet.row(row_idx)
                row_text = [str(cell.value) for cell in row if cell.value]
                if row_text:
                    all_text.append(" ".join(row_text))
        return "\n".join(all_text)
    except Exception as e:
        return f"Error reading XLS: {str(e)}"
 
def read_file_content(file_path):
    ext = file_path.split('.')[-1].lower()
    if ext == "txt":
        return read_txt(file_path)
    elif ext == "pdf":
        return read_pdf(file_path)
    elif ext == "docx":
        return read_docx(file_path)
    elif ext == "xlsx":
        return read_xlsx(file_path)
    elif ext == "xls":
        return read_xls(file_path)
    return ""
 
@app.route("/chat/respond", methods=["POST"])
def chat_respond():
    payload = request.json
    message = payload.get("message", "").strip()
    if not message:
        return jsonify({"response": "Oops! Type something first, love!"}), 400
 
    all_text = ""
    for root, dirs, files in os.walk(UPLOAD_DIR):
        for name in files:
            file_path = os.path.join(root, name)
            if os.path.isfile(file_path):
                all_text += "\n" + read_file_content(file_path)
 
    if not all_text.strip():
        return jsonify({"response": "Hey, I couldn't find any readable content in the uploaded files 😔"})
 
    lines = [line.strip() for line in all_text.splitlines() if line.strip()]
    best_line = ""
    best_score = 0
 
    for line in lines:
        score = fuzz.partial_ratio(message.lower(), line.lower())
        if score > best_score:
            best_score = score
            best_line = line
 
    if best_score > 50:
        response = f"Hey! I found this for you 😊: \"{best_line}\""
    else:
        response = "Hmm... I tried, but couldn't find a clear answer in the files 😅"
 
    return jsonify({"response": response})
 
if __name__ == "__main__":
    app.run(debug=True)